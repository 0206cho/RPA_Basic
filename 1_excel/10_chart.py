from openpyxl import load_workbook
wb = load_workbook("sample_5.xlsx")
ws = wb.active

# 차트 사용하기 위해 import, BarChart사용. Reference 함께 추가
from openpyxl.chart import BarChart, Reference, LineChart

# B2:C11 까지의 데이터를 차트로 생성
# 어떤 데이터 범위를 쓸지 정의. (현재워크시트 정보, 범위)
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # 차트 종류 설정(Bar, Line, Pie,...,)
# bar_chart.add_data(bar_value) # 차트 데이터 추가
# ws.add_chart(bar_chart, "E1") #차트 넣을 위치 정의

# LineChart - 제목포함한 B1:C11까지의 데이터
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) # titles_from_data=True : 계열부분 제목에서 가져옴
line_chart.title = "성적표" #제목
line_chart.style = 10 # 미리 정의된 스타일을 적용, 사용자가 개별 지정도 가능
line_chart.y_axis.title = "점수" #Y축의 제목
line_chart.x_axis.title = "번호" #X축의 제목
ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")